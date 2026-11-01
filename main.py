import re
from typing import NamedTuple

import qrcode
import zxingcpp
import openpyxl
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps, ImageWin

Format = zxingcpp.BarcodeFormat


class BarcodeType(NamedTuple):
    format: Format | None  # None means QR Code, made by generate_qr
    layout: str  # "linear" (bars + text, adjustable height), "stacked" or "matrix"
    example: str
    hint: str
    options: dict = {}  # extra zxingcpp.create_barcode arguments


QR_CODE = "QR Code"

BARCODE_TYPES = {
    # 2D codes
    QR_CODE: BarcodeType(None, "matrix", "https://example.com",
                         "Any text or link. Phones scan it with the camera."),
    "Data Matrix": BarcodeType(Format.DataMatrix, "matrix", "SN-2026-000123",
                               "Any text. Small square code on electronics, medicine and parcels.",
                               {"forceSquare": True}),
    "PDF417": BarcodeType(Format.PDF417, "stacked", "ATH-LHR 2026-11-01 Seat 14C",
                          "Any text, even long. ID cards, boarding passes, shipping labels."),
    "Aztec": BarcodeType(Format.Aztec, "matrix", "TICKET-123456",
                         "Any text. Train and plane tickets."),
    # Retail
    "EAN-13": BarcodeType(Format.EAN13, "linear", "590123412345",
                          "12 digits (the 13th check digit is added). Products in Europe."),
    "EAN-8": BarcodeType(Format.EAN8, "linear", "9638507",
                         "7 digits (check digit added). Small packages."),
    "UPC-A": BarcodeType(Format.UPCA, "linear", "03600029145",
                         "11 digits (check digit added). Products in North America."),
    "UPC-E": BarcodeType(Format.UPCE, "linear", "0425261",
                         "7 digits starting with 0 or 1. Compressed UPC for small packages."),
    "ISBN": BarcodeType(Format.ISBN, "linear", "9783161484100",
                        "10 or 13-digit ISBN. Books."),
    # Logistics (GS1)
    "GS1-128": BarcodeType(Format.Code128, "linear", "(01)09501101530003(10)ABC123",
                           "GS1 data with application identifiers in brackets. Shipping labels.",
                           {"gs1": True}),
    "ITF-14": BarcodeType(Format.ITF14, "linear", "0950110153000",
                          "13 digits (check digit added). Outer cartons and pallets."),
    "GS1 DataBar": BarcodeType(Format.DataBarOmni, "linear", "0950110153000",
                               "13-digit GTIN (check digit added). Fresh produce and coupons."),
    "GS1 DataBar Expanded": BarcodeType(Format.DataBarExp, "linear", "(01)09501101530003(3103)000123",
                                        "GS1 data with application identifiers in brackets. Weighed food."),
    # General purpose
    "Code 128": BarcodeType(Format.Code128, "linear", "ABC-12345",
                            "Latin letters, digits and symbols. The most common general barcode."),
    "Code 39": BarcodeType(Format.Code39, "linear", "CODE-39",
                           "Capital letters, digits, space and - . $ / + %. Industry and ID badges."),
    "Code 93": BarcodeType(Format.Code93, "linear", "CODE-93",
                           "Like Code 39 but more compact."),
    "Interleaved 2 of 5": BarcodeType(Format.ITF, "linear", "12345678",
                                      "Digits only. Warehouses and distribution."),
    "Codabar": BarcodeType(Format.Codabar, "linear", "A123456B",
                           "Digits and - $ : / . +, starting and ending with A, B, C or D. Libraries, blood banks."),
    "PZN": BarcodeType(Format.PZN, "linear", "123456",
                       "6 or 7 digits. German pharmaceuticals."),
}


def generate_qr(data, box_size=10):
    qr = qrcode.QRCode(box_size=box_size)
    qr.add_data(data)

    return qr.make_image()


def generate_barcode(data, barcode_type, scale=2, bar_height=None):
    """scale is pixels per module; bar_height (pixels) applies to linear barcodes only.

    Raises ValueError with a readable message when the data does not fit the type.
    """
    kind = BARCODE_TYPES[barcode_type]
    try:
        code = zxingcpp.create_barcode(data, kind.format, **kind.options)
    except ValueError as e:
        raise ValueError(re.sub(r"^Error \d+: |\s*\(retval: \d+\)$", "", str(e))) from None

    pixels = memoryview(code.to_image(scale=scale, add_hrt=kind.layout == "linear"))
    height, width = pixels.shape
    image = Image.frombytes("L", (width, height), pixels.tobytes())

    if kind.layout == "linear":
        if bar_height:
            image = _set_bar_height(image, bar_height)
        # zxing only adds quiet zones left and right of linear barcodes
        image = ImageOps.expand(image, border=(0, 5 * scale), fill=255)
    return image


def _set_bar_height(image, bar_height):
    """Stretch or shrink the bars of a linear barcode, keeping the text below them.

    The bars are the longest run of identical rows that mix black and white, which skips
    blank margins and the solid bearer bars of ITF-14.
    """
    data, width = image.tobytes(), image.width
    rows = [data[y * width:(y + 1) * width] for y in range(image.height)]
    top = bottom = 0
    y = 0
    while y < len(rows):
        end = y + 1
        while end < len(rows) and rows[end] == rows[y]:
            end += 1
        if 0 in rows[y] and 255 in rows[y] and end - y > bottom - top:
            top, bottom = y, end
        y = end

    result = Image.new("L", (width, image.height - (bottom - top) + bar_height), 255)
    result.paste(image.crop((0, 0, width, top)), (0, 0))
    result.paste(image.crop((0, top, width, top + 1)).resize((width, bar_height)), (0, top))
    result.paste(image.crop((0, bottom, width, image.height)), (0, top + bar_height))
    return result


def read_excel(file_path):
    """Return the first sheet as {column letter: [cell text, ...]}, top row first."""
    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    workbook.close()

    width = max((len(row) for row in rows), default=0)
    return {
        get_column_letter(i + 1): [_cell_text(row[i] if i < len(row) else None) for row in rows]
        for i in range(width)
    }


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)

    return str(value).strip()


MAX_SEQUENCE = 10000


def make_sequence(start, end, step=1, prefix="", suffix="", digits=0):
    """Return e.g. ["INV-0001", "INV-0002", ...]; counts down when start > end."""
    if step < 1:
        raise ValueError("Step must be 1 or more.")
    if abs(end - start) // step + 1 > MAX_SEQUENCE:
        raise ValueError(f"That is more than {MAX_SEQUENCE} values. Use a smaller range or a bigger step.")

    direction = 1 if end >= start else -1
    return [f"{prefix}{n:0{digits}d}{suffix}" for n in range(start, end + direction, step * direction)]


def list_printers():
    """Return (installed printer names, default printer name). Windows only."""
    import win32print

    flags = win32print.PRINTER_ENUM_LOCAL | win32print.PRINTER_ENUM_CONNECTIONS
    return [printer[2] for printer in win32print.EnumPrinters(flags)], win32print.GetDefaultPrinter()


def print_images(images, printer, columns=1, rows=1, title="Barcodes", output_file=None):
    """Print images on printer, columns x rows per page, using the page size set in the printer driver.

    Each image is enlarged by a whole number so bars stay exact; pass images rendered at
    one pixel per module. output_file is for printers that write files (e.g. Print to PDF).
    """
    import win32con
    import win32ui

    dc = win32ui.CreateDC()
    dc.CreatePrinterDC(printer)
    try:
        cell_width = dc.GetDeviceCaps(win32con.HORZRES) // columns
        cell_height = dc.GetDeviceCaps(win32con.VERTRES) // rows
        margin = min(cell_width, cell_height) // 20
        per_page = columns * rows

        dc.StartDoc(title, output_file) if output_file else dc.StartDoc(title)
        for first in range(0, len(images), per_page):
            dc.StartPage()
            for i, image in enumerate(images[first:first + per_page]):
                room_width, room_height = cell_width - 2 * margin, cell_height - 2 * margin
                factor = min(room_width // image.width, room_height // image.height)
                if factor >= 1:
                    size = (image.width * factor, image.height * factor)
                else:  # bigger than the label even at 1px per module: shrink to fit
                    ratio = min(room_width / image.width, room_height / image.height)
                    size = (max(1, int(image.width * ratio)), max(1, int(image.height * ratio)))

                x = (i % columns) * cell_width + (cell_width - size[0]) // 2
                y = (i // columns) * cell_height + (cell_height - size[1]) // 2
                dib = ImageWin.Dib(image.convert("RGB").resize(size, Image.NEAREST))
                dib.draw(dc.GetHandleOutput(), (x, y, x + size[0], y + size[1]))
            dc.EndPage()
        dc.EndDoc()
    finally:
        dc.DeleteDC()


if __name__ == "__main__":
    url = input("Enter the URL: ").strip()

    file_path ="qrcode.png"

    img = generate_qr(url)
    img.save(file_path)

    print("QR Code was generated !")
